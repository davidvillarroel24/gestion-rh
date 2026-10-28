import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from core import roles
from proyectos.models import Avance, Desembolso, EstadoProyecto, Financiamiento, Indicador, Medicion
from proyectos.tests import Base

from . import servicios

HOY = date(2026, 10, 28)


class ReportesBase(Base):
    def setUp(self):
        super().setUp()
        # P1: mancomunado Punata + Tacachi, 1000 (800 GADC + 200 GAM Punata); P2: solo Punata, 500 (GADC)
        self.p1 = self.proyecto(nombre="Mancomunado", municipios=[self.m1, self.m2], gestion=2026)
        self.p2 = self.proyecto(nombre="Solo Punata", municipios=[self.m1], gestion=2025)
        self.f1 = Financiamiento.objects.create(proyecto=self.p1, entidad=self.gadc, monto_comprometido=Decimal("800"))
        Financiamiento.objects.create(proyecto=self.p1, entidad=self.gam, monto_comprometido=Decimal("200"))
        self.f2 = Financiamiento.objects.create(proyecto=self.p2, entidad=self.gadc, monto_comprometido=Decimal("500"))
        Desembolso.objects.create(financiamiento=self.f1, fecha=date(2026, 3, 10), monto=Decimal("400"))
        Desembolso.objects.create(financiamiento=self.f2, fecha=date(2026, 5, 20), monto=Decimal("100"))
        self.tec = self.usuario("tec", roles.TECNICO, unidades=[self.riego])

    def lista(self, usuario=None, **filtros):
        return servicios.proyectos(usuario or self.tec, filtros)


class AgrupadoTests(ReportesBase):
    def test_por_municipio_no_duplica_el_total(self):
        r = servicios.agrupar(self.lista(), "municipio")
        filas = {f["grupo"]: f for f in r["filas"]}
        self.assertEqual((filas["Punata"]["proyectos"], filas["Punata"]["comprometido"]), (2, 1500.0))
        self.assertEqual((filas["Tacachi"]["proyectos"], filas["Tacachi"]["comprometido"]), (1, 1000.0))
        self.assertEqual(r["total"]["comprometido"], 1500.0, "el total cuenta cada proyecto una vez")
        self.assertEqual(r["total"]["proyectos"], 2)
        self.assertIn("aparece en cada uno", r["nota"])
        self.assertEqual(filas["Punata"]["aporte_municipal"], 200.0, "contraparte del GAM Punata")
        self.assertEqual(filas["Tacachi"]["aporte_municipal"], 0.0)

    def test_por_unidad_y_gestion_suman_el_total(self):
        for por in ("unidad", "gestion", "estado"):
            r = servicios.agrupar(self.lista(), por)
            self.assertEqual(sum(f["comprometido"] for f in r["filas"]), r["total"]["comprometido"], por)
            self.assertIsNone(r["nota"], por)
        gestiones = {f["grupo"]: f["comprometido"] for f in servicios.agrupar(self.lista(), "gestion")["filas"]}
        self.assertEqual(gestiones, {"2026": 1000.0, "2025": 500.0})

    def test_por_fuente_y_entidad(self):
        fuentes = {f["grupo"]: f for f in servicios.agrupar(self.lista(), "fuente")["filas"]}
        self.assertEqual((fuentes["Gobernación"]["comprometido"], fuentes["Gobernación"]["desembolsado"]), (1300.0, 500.0))
        self.assertEqual((fuentes["Municipio"]["comprometido"], fuentes["Municipio"]["proyectos"]), (200.0, 1))
        self.assertEqual(sum(f["comprometido"] for f in fuentes.values()), 1500.0)
        entidades = {f["grupo"] for f in servicios.agrupar(self.lista(), "entidad")["filas"]}
        self.assertEqual(entidades, {"GADC", "GAM Punata"})

    def test_por_region_y_provincia(self):
        self.assertEqual([f["grupo"] for f in servicios.agrupar(self.lista(), "region")["filas"]], ["Valles"])
        self.assertEqual(servicios.agrupar(self.lista(), "provincia")["filas"][0]["proyectos"], 2)


class ResumenYCurvaTests(ReportesBase):
    def test_tarjetas(self):
        Avance.objects.create(proyecto=self.p1, fecha=date(2026, 4, 1), fisico_pct=Decimal("50"))
        Avance.objects.create(proyecto=self.p2, fecha=date(2026, 6, 1), fisico_pct=Decimal("20"))
        t = servicios.resumen(self.lista())
        self.assertEqual((t["proyectos"], t["comprometido"], t["desembolsado"], t["saldo"]), (2, 1500.0, 500.0, 1000.0))
        self.assertEqual(t["avance_financiero"], 33.3)
        self.assertEqual(t["avance_fisico"], 40.0, "ponderado por monto: (50·1000 + 20·500) / 1500")
        self.assertEqual(t["municipios"], 2)

    def test_curva_s(self):
        Avance.objects.create(proyecto=self.p1, fecha=date(2026, 4, 1), fisico_pct=Decimal("60"))
        c = servicios.curva_s(self.lista(), hoy=date(2026, 6, 15))
        self.assertEqual(c["meses"], ["2026-03", "2026-04", "2026-05", "2026-06"])
        self.assertEqual(c["financiero"], [26.7, 26.7, 33.3, 33.3], "400 de 1500 en marzo; +100 en mayo")
        self.assertEqual(c["fisico"], [0.0, 40.0, 40.0, 40.0], "60 % de un proyecto que pesa 1000 de 1500")
        self.assertEqual(servicios.curva_s([], hoy=HOY)["meses"], [])

    def test_indicadores_antes_y_despues(self):
        prod = Indicador.objects.create(nombre="Producción de papa", unidad_medida="t/año")
        for p, base, despues in ((self.p1, 10, 18), (self.p2, 5, 6)):
            Medicion.objects.create(proyecto=p, indicador=prod, momento=Medicion.LINEA_BASE, fecha=date(2025, 1, 1), valor=base)
            Medicion.objects.create(proyecto=p, indicador=prod, momento=Medicion.SEGUIMIENTO, fecha=date(2026, 1, 1), valor=despues - 1)
            Medicion.objects.create(proyecto=p, indicador=prod, momento=Medicion.SEGUIMIENTO, fecha=date(2026, 9, 1), valor=despues)
        fila = servicios.indicadores(self.lista())[0]
        self.assertEqual((fila["linea_base"], fila["seguimiento"], fila["cambio"], fila["proyectos"]), (15.0, 24.0, 9.0, 2),
                         "se suma el ÚLTIMO seguimiento de cada proyecto")


class AlertasTests(ReportesBase):
    def tipos(self, proyecto):
        return {a["tipo"] for a in servicios.alertas(self.lista(), hoy=HOY) if a["proyecto_id"] == proyecto.pk}

    def test_alertas_de_proyecto_en_curso(self):
        self.p1.fecha_orden_proceder, self.p1.fecha_conclusion = date(2026, 1, 1), date(2026, 9, 30)
        self.p1.save()
        # 40 % pagado (400 de 1000) y ningún avance físico cargado.
        self.assertEqual(self.tipos(self.p1), {"Plazo vencido", "Sin avance registrado", "Pagado por delante de lo ejecutado"})
        Avance.objects.create(proyecto=self.p1, fecha=date(2026, 5, 1), fisico_pct=Decimal("30"))
        self.assertEqual(self.tipos(self.p1), {"Plazo vencido", "Sin avance reciente"}, "brecha 40 − 30 < 25")
        Desembolso.objects.create(financiamiento=self.f1, fecha=date(2026, 6, 1), monto=Decimal("300"))  # 70 % vs 30 %
        self.assertIn("Pagado por delante de lo ejecutado", self.tipos(self.p1))

    def test_proyecto_cerrado_solo_pide_seguimiento(self):
        concluido = EstadoProyecto.objects.create(nombre="Concluido", es_final=True)
        from core.models import TipoEvento
        from proyectos.models import Evento
        Evento.objects.create(proyecto=self.p2, tipo=TipoEvento.objects.create(nombre="Cierre"), fecha=date(2026, 1, 1),
                              descripcion="Fin", estado_nuevo=concluido)
        self.p2.fecha_conclusion = date(2025, 12, 1)
        self.p2.save()
        self.assertEqual(self.tipos(self.p2), {"Cerrado sin medición de seguimiento"})

    def test_desembolsado_mayor_a_comprometido(self):
        Desembolso.objects.create(financiamiento=self.f2, fecha=date(2026, 7, 1), monto=Decimal("450"), excede_confirmado=True)
        self.assertIn("Desembolsado mayor a lo comprometido", self.tipos(self.p2))


class VistasTests(ReportesBase):
    def test_pagina_y_api(self):
        self.assertEqual(self.client.get("/reportes/").status_code, 302, "pide login")
        self.client.force_login(self.tec)
        self.assertContains(self.client.get("/reportes/"), "reportes/reportes")
        d = self.client.get("/reportes/api/", {"por": "fuente"}).json()
        self.assertEqual(set(d), {"tarjetas", "agrupado", "curva_s", "indicadores", "alertas"})
        self.assertEqual(d["agrupado"]["por"], "fuente")
        self.assertEqual(self.client.get("/reportes/api/", {"por": "x", "unidad": "abc"}).json()["agrupado"]["por"], "municipio",
                         "parámetros inválidos se ignoran")

    def test_filtros_de_region_provincia_y_alcance(self):
        self.client.force_login(self.tec)
        self.assertEqual(self.client.get("/reportes/api/", {"municipio": self.m2.pk}).json()["tarjetas"]["proyectos"], 1)
        self.assertEqual(self.client.get("/reportes/api/", {"region": self.m1.region_id}).json()["tarjetas"]["proyectos"], 2)
        self.client.force_login(self.usuario("enc", roles.ENCARGADO_MUNICIPAL, municipios=[self.m2]))
        self.assertEqual(self.client.get("/reportes/api/").json()["tarjetas"]["proyectos"], 1, "solo su municipio")

    def test_excel(self):
        self.client.force_login(self.tec)
        r = self.client.get("/reportes/excel/", {"por": "municipio"})
        libro = load_workbook(io.BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ["Por Municipio", "Resumen", "Indicadores", "Alertas"])
        filas = list(libro["Por Municipio"].iter_rows(values_only=True))
        self.assertEqual(filas[0][-1], "Aporte municipal (Bs)")
        total = next(f for f in filas if f[0] == "TOTAL (proyectos distintos)")
        self.assertEqual(total[2], 1500)
