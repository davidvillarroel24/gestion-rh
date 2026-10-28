import io
from datetime import date

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core import alcance
from proyectos.models import EstadoProyecto
from territorio.models import Provincia, Region

from . import servicios

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _por(request):
    por = request.GET.get("por", "municipio")
    return por if por in servicios.DIMENSIONES else "municipio"


@login_required
def pagina(request):
    u = request.user
    return render(request, "reportes/reportes.html", {
        "filtros": {
            "unidades": [{"id": x.id, "nombre": x.nombre} for x in alcance.unidades_visibles(u).filter(activo=True)],
            "estados": [{"id": x.id, "nombre": x.nombre} for x in EstadoProyecto.objects.filter(activo=True)],
            "regiones": [{"id": x.id, "nombre": x.nombre} for x in Region.objects.order_by("nombre")],
            "provincias": [{"id": x.id, "nombre": x.nombre} for x in Provincia.objects.order_by("nombre")],
            "municipios": [{"id": m.id, "nombre": m.nombre, "provincia": m.provincia_id, "region": m.region_id}
                           for m in alcance.municipios_visibles(u).order_by("nombre")],
        },
        "dimensiones": [(k, v[0]) for k, v in servicios.DIMENSIONES.items()],
        "medidas": list(servicios.MEDIDAS.items()),
    })


@login_required
def datos(request):
    """Todo lo que muestra la página de reportes, con los filtros de la URL."""
    if not request.user.has_perm("proyectos.view_proyecto"):
        return JsonResponse({"error": "Sin permiso"}, status=403)
    lista = servicios.proyectos(request.user, request.GET)
    return JsonResponse({
        "tarjetas": servicios.resumen(lista),
        "agrupado": servicios.agrupar(lista, _por(request)),
        "curva_s": servicios.curva_s(lista),
        "indicadores": servicios.indicadores(lista),
        "alertas": servicios.alertas(lista),
    })


def _hoja(libro, titulo, encabezados, filas):
    hoja = libro.create_sheet(titulo)
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="D1E7DD")
    for fila in filas:
        hoja.append(fila)
    for i, encabezado in enumerate(encabezados, start=1):
        hoja.column_dimensions[hoja.cell(row=1, column=i).column_letter].width = max(14, min(60, len(encabezado) + 6))
    return hoja


@login_required
def excel(request):
    """Reporte descargable: agrupación elegida, resumen, indicadores y alertas."""
    if not request.user.has_perm("proyectos.view_proyecto"):
        return HttpResponse(status=403)
    lista = servicios.proyectos(request.user, request.GET)
    agrupado = servicios.agrupar(lista, _por(request))
    tarjetas = servicios.resumen(lista)
    libro = Workbook()
    libro.remove(libro.active)

    con_aporte = agrupado["por"] == "municipio"
    encabezados = [agrupado["dimension"], "Proyectos", "Comprometido (Bs)", "Desembolsado (Bs)", "Avance financiero (%)"]
    filas = [[f["grupo"], f["proyectos"], f["comprometido"], f["desembolsado"], f["avance_financiero"]]
             + ([f.get("aporte_municipal", 0)] if con_aporte else []) for f in agrupado["filas"]]
    t = agrupado["total"]
    filas.append(["TOTAL (proyectos distintos)", t["proyectos"], t["comprometido"], t["desembolsado"], t["avance_financiero"]])
    hoja = _hoja(libro, f"Por {agrupado['dimension']}"[:31], encabezados + (["Aporte municipal (Bs)"] if con_aporte else []), filas)
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)
    if agrupado["nota"]:
        hoja.append([])
        hoja.append([agrupado["nota"]])

    _hoja(libro, "Resumen", ["Dato", "Valor"], [
        ["Proyectos", tarjetas["proyectos"]], ["Monto comprometido (Bs)", tarjetas["comprometido"]],
        ["Desembolsado (Bs)", tarjetas["desembolsado"]], ["Saldo por desembolsar (Bs)", tarjetas["saldo"]],
        ["Avance financiero (%)", tarjetas["avance_financiero"]], ["Avance físico ponderado (%)", tarjetas["avance_fisico"]],
        ["Municipios abarcados", tarjetas["municipios"]], ["Proyectos cerrados", tarjetas["concluidos"]],
    ])
    _hoja(libro, "Indicadores", ["Indicador", "Desagregación", "Unidad", "Proyectos", "Línea base", "Meta", "Seguimiento", "Cambio"],
          [[i["indicador"], i["desagregacion"], i["unidad_medida"], i["proyectos"], i["linea_base"], i["meta"],
            i["seguimiento"], i["cambio"]] for i in servicios.indicadores(lista)])
    _hoja(libro, "Alertas", ["Nivel", "Alerta", "Código", "Proyecto", "Estado", "Detalle"],
          [[a["nivel"], a["tipo"], a["codigo"], a["nombre"], a["estado"], a["detalle"]] for a in servicios.alertas(lista)])

    salida = io.BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(salida.getvalue(), content_type=XLSX)
    respuesta["Content-Disposition"] = f'attachment; filename="reporte-{agrupado["por"]}-{date.today():%Y-%m-%d}.xlsx"'
    return respuesta
