"""
Importación y exportación de planillas (Excel / CSV) según importador/esquemas.py.

Validar y aplicar usan exactamente el mismo código: todo corre dentro de una
transacción que se deshace al validar (o si hubo algún error). Así lo que se
muestra como "listo para confirmar" es exactamente lo que se va a guardar.
"""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.models import Entidad, TipoFuente
from core.texto import normalizar
from territorio.models import Municipio

from . import valores as v
from .esquemas import TABLAS, Contexto, ErrorFila, convertir
from .models import Cambio, Importacion

FILAS_CON_LISTA = 2000
ROJO = PatternFill("solid", fgColor="F8D7DA")
GRIS = PatternFill("solid", fgColor="E9ECEF")
VERDE = PatternFill("solid", fgColor="D1E7DD")


class ErrorArchivo(Exception):
    pass


# ---------------------------------------------------------------- lectura


def _filas_crudas(archivo, nombre, tabla):
    nombre = nombre.lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        try:
            libro = load_workbook(archivo, read_only=True, data_only=True)
        except Exception:
            raise ErrorArchivo("No se pudo leer el Excel (¿está dañado o es otro formato?).")
        hojas = {normalizar(h): h for h in libro.sheetnames}
        hoja = libro[hojas.get(normalizar(tabla.titulo)) or hojas.get("datos") or libro.sheetnames[0]]
        return [list(f) for f in hoja.iter_rows(values_only=True)]
    if nombre.endswith(".csv"):
        crudo = archivo.read()
        for codificacion in ("utf-8-sig", "cp1252"):
            try:
                texto = crudo.decode(codificacion)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        except csv.Error:
            dialecto = csv.excel
        return list(csv.reader(io.StringIO(texto), dialecto))
    raise ErrorArchivo("Formato no soportado: use .xlsx o .csv.")


def leer(archivo, nombre, tabla):
    """Devuelve [(número de fila en el archivo, {clave_columna: valor})]."""
    filas = _filas_crudas(archivo, nombre, tabla)
    por_titulo = {}
    for c in tabla.columnas_importables:
        por_titulo[normalizar(c.titulo)] = c
        por_titulo[normalizar(c.clave)] = c

    # Encabezados: la primera fila que nombra alguna columna conocida ("Unidad*" = "unidad").
    for encabezado, fila in enumerate(filas):
        indices = {j: por_titulo[normalizar(celda)] for j, celda in enumerate(fila)
                   if not v.vacio(celda) and normalizar(celda) in por_titulo}
        if indices:
            break
    else:
        raise ErrorArchivo(f"No se encontró la fila de encabezados de «{tabla.titulo}». Use la plantilla de esta tabla.")
    faltan = [c.titulo for c in tabla.columnas_importables if c.obligatorio and c not in indices.values()]
    if faltan:
        raise ErrorArchivo(f"Faltan columnas obligatorias: {', '.join(faltan)}.")

    resultado = []
    for numero, fila in enumerate(filas[encabezado + 1:], start=encabezado + 2):
        celdas = {col.clave: fila[j] for j, col in indices.items() if j < len(fila)}
        if all(v.vacio(x) for x in celdas.values()):
            continue
        resultado.append((numero, celdas))
    return resultado


# ---------------------------------------------------------------- proceso


def procesar_filas(tabla, filas, ctx):
    resultado = {"filas": len(filas), "creados": 0, "actualizados": 0, "relaciones": 0, "sin_cambios": 0, "errores": []}
    for numero, celdas in filas:
        errores = []
        cambios_antes = len(ctx.cambios)
        try:
            with transaction.atomic():  # cada fila es todo o nada (incluye catálogos creados al vuelo)
                datos = {}
                for col in tabla.columnas_importables:
                    valor = celdas.get(col.clave)
                    if v.vacio(valor):
                        if col.obligatorio:
                            errores.append(f"{col.titulo}: es obligatorio.")
                        datos[col.clave] = None
                        continue
                    try:
                        datos[col.clave] = convertir(col, valor, datos, ctx)
                    except v.ErrorValor as e:
                        errores.append(f"{col.titulo}: {e}")
                        datos[col.clave] = None
                    except ValidationError as e:
                        errores.append(f"{col.titulo}: {' '.join(e.messages)}")
                        datos[col.clave] = None
                if errores:
                    raise ErrorFila(errores)
                accion = tabla.guardar(datos, ctx)
        except ErrorFila as e:
            del ctx.cambios[cambios_antes:]
            resultado["errores"].append({"fila": numero, "mensajes": e.mensajes})
            continue
        clave = {"creado": "creados", "actualizado": "actualizados", "relacion": "relaciones"}.get(accion, "sin_cambios")
        resultado[clave] += 1
    return resultado


def _guardar_cambios(importacion, ctx):
    tipos = {}
    for c in ctx.cambios:
        modelo = type(c["objeto"])
        tipos.setdefault(modelo, ContentType.objects.get_for_model(modelo))
        Cambio.objects.create(importacion=importacion, tipo_contenido=tipos[modelo], objeto_id=c["objeto"].pk,
                              accion=c["accion"], version_previa=c["previa"], version_importada=c["importada"],
                              extra=c["extra"])


def ejecutar(importacion, aplicar=False):
    """Valida (aplicar=False) o aplica una importación tabular y actualiza su estado y resumen."""
    tabla = TABLAS[importacion.destino]
    ctx = Contexto(importacion.usuario, crear_catalogos=importacion.opciones.get("crear_catalogos", False))
    try:
        with importacion.archivo.open("rb") as archivo:
            filas = leer(archivo, importacion.nombre_archivo, tabla)
    except ErrorArchivo as e:
        importacion.resumen = {"filas": 0, "errores": [{"fila": None, "mensajes": [str(e)]}]}
        importacion.estado = Importacion.CON_ERRORES
        importacion.save()
        return importacion

    with transaction.atomic():
        resultado = procesar_filas(tabla, filas, ctx)
        if resultado["errores"] or not aplicar:
            transaction.set_rollback(True)
        else:
            _guardar_cambios(importacion, ctx)

    importacion.resumen = resultado
    if resultado["errores"]:
        importacion.estado = Importacion.CON_ERRORES
    elif aplicar:
        importacion.estado = Importacion.APLICADA
        importacion.aplicada_en = timezone.now()
    else:
        importacion.estado = Importacion.VALIDADA
    importacion.save()
    return importacion


# ---------------------------------------------------------------- plantillas y exportación


def _formatear(valor):
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def _encabezado(hoja, columnas, marcar_obligatorias=True):
    for j, col in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=j, value=col.titulo + ("*" if marcar_obligatorias and col.obligatorio else ""))
        celda.font = Font(bold=True)
        celda.fill = GRIS if col.solo_exportar else VERDE
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        if col.ayuda or col.solo_exportar:
            celda.comment = Comment(col.ayuda or "Calculado: se ignora al importar.", "Gestión RH")
        hoja.column_dimensions[get_column_letter(j)].width = max(14, min(40, len(col.titulo) + 4))
    hoja.freeze_panes = "A2"


def _valores_lista(col, ctx):
    ref = col.referencia
    if col.tipo == "opcion":
        return [etiqueta for _, etiqueta in col.opciones], True
    if ref == "proyecto":
        return list(ctx.proyectos.order_by("codigo").values_list("codigo", flat=True)), True
    if ref == "unidad":
        return list(ctx.unidades.filter(activo=True).values_list("nombre", flat=True)), True
    if ref == "municipio":
        return list(Municipio.objects.order_by("nombre").values_list("nombre", flat=True)), True
    if ref == "entidad":
        return list(Entidad.objects.filter(activo=True).order_by("nombre").values_list("nombre", flat=True)), False
    if ref == "tipo_fuente":
        return list(TipoFuente.objects.filter(activo=True).values_list("nombre", flat=True)), True
    if ref is not None:
        qs = ref.objects.filter(activo=True)
        if ref.campos_ambito:
            qs = qs.filter(unidad__in=ctx.unidades)
        return sorted(set(qs.values_list("nombre", flat=True))), False
    return None, False


def plantilla(tabla, usuario):
    """Libro con hojas Datos (encabezados + listas desplegables), Instrucciones y Listas."""
    ctx = Contexto(usuario)
    libro = Workbook()
    datos = libro.active
    datos.title = "Datos"
    columnas = tabla.columnas_importables
    _encabezado(datos, columnas)

    listas = libro.create_sheet("Listas")
    columna_lista = 0
    for j, col in enumerate(columnas, start=1):
        valores, estricta = _valores_lista(col, ctx)
        if not valores:
            continue
        columna_lista += 1
        letra = get_column_letter(columna_lista)
        listas.cell(row=1, column=columna_lista, value=col.titulo).font = Font(bold=True)
        for i, valor in enumerate(valores, start=2):
            listas.cell(row=i, column=columna_lista, value=valor)
        validacion = DataValidation(type="list", formula1=f"=Listas!${letra}$2:${letra}${len(valores) + 1}",
                                    allow_blank=True, showErrorMessage=True,
                                    errorStyle="stop" if estricta else "warning",
                                    error="Elija un valor de la lista." if estricta else
                                    "No está en la lista: se creará solo si marca «crear valores de catálogo nuevos».")
        datos.add_data_validation(validacion)
        destino = get_column_letter(j)
        validacion.add(f"{destino}2:{destino}{FILAS_CON_LISTA}")

    instrucciones = libro.create_sheet("Instrucciones", 0)
    instrucciones.append([f"Plantilla: {tabla.titulo}"])
    instrucciones["A1"].font = Font(bold=True, size=14)
    instrucciones.append([tabla.descripcion])
    instrucciones.append(["Llene la hoja «Datos» desde la fila 2. Las columnas con * son obligatorias. "
                          "Las celdas con lista desplegable toman sus valores de la hoja «Listas»."])
    instrucciones.append([])
    instrucciones.append(["Columna", "Obligatoria", "Tipo", "Ayuda / valores", "Ejemplo"])
    for celda in instrucciones[5]:
        celda.font = Font(bold=True)
        celda.fill = VERDE
    tipos = {"texto": "Texto", "entero": "Número entero", "decimal": "Número", "fecha": "Fecha (dd/mm/aaaa)",
             "booleano": "Sí / No", "opcion": "Opción de la lista", "referencia": "Nombre de la lista"}
    for i, col in enumerate(columnas):
        ayuda = col.ayuda or (f"Opciones: {', '.join(e for _, e in col.opciones)}" if col.opciones else "")
        ejemplo = tabla.ejemplo[i] if i < len(tabla.ejemplo) else ""
        instrucciones.append([col.titulo, "Sí" if col.obligatorio else "", tipos[col.tipo], ayuda, ejemplo])
    for letra, ancho in zip("ABCDE", (32, 12, 20, 70, 30)):
        instrucciones.column_dimensions[letra].width = ancho
    libro.active = 1
    return libro


def exportar(tablas, proyectos, usuario):
    """Libro con una hoja por tabla, con los nombres resueltos (nunca IDs)."""
    libro = Workbook()
    libro.remove(libro.active)
    for tabla in tablas:
        hoja = libro.create_sheet(tabla.titulo[:31])
        _encabezado(hoja, tabla.columnas, marcar_obligatorias=False)
        for fila in tabla.exportar(proyectos, usuario):
            hoja.append([_formatear(x) for x in fila])
        for j, col in enumerate(tabla.columnas, start=1):
            if col.tipo == "fecha":
                for celda in hoja.iter_rows(min_row=2, min_col=j, max_col=j):
                    celda[0].number_format = "DD/MM/YYYY"
            elif col.tipo == "decimal":
                for celda in hoja.iter_rows(min_row=2, min_col=j, max_col=j):
                    celda[0].number_format = "#,##0.00"
    return libro


def libro_de_errores(importacion):
    """El archivo original con una columna «Errores» y las filas con problemas en rojo."""
    tabla = TABLAS[importacion.destino]
    with importacion.archivo.open("rb") as archivo:
        filas = _filas_crudas(archivo, importacion.nombre_archivo, tabla)
    errores = {e["fila"]: " | ".join(e["mensajes"]) for e in importacion.resumen.get("errores", []) if e.get("fila")}
    libro = Workbook()
    hoja = libro.active
    hoja.title = tabla.titulo[:31]
    for numero, fila in enumerate(filas, start=1):
        valores = [x.strftime("%d/%m/%Y") if isinstance(x, (date, datetime)) else x for x in fila]
        hoja.append(valores + [errores.get(numero, "Errores" if numero == 1 else "")])
        if numero in errores:
            for celda in hoja[numero]:
                celda.fill = ROJO
    return libro
