import io
import json
import zipfile
from datetime import date
from decimal import Decimal

from django.contrib.gis.geos import Point
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from openpyxl import Workbook, load_workbook

from core import roles
from core.models import Entidad, TipoEvento, Usuario
from proyectos.models import Desembolso, Evento, Financiamiento, Geometria, Indicador, Medicion, Proyecto
from proyectos.tests import Base
from territorio.models import Municipio

from .models import Importacion


def xlsx(encabezados, *filas, hoja="Datos"):
    libro = Workbook()
    h = libro.active
    h.title = hoja
    h.append(encabezados)
    for fila in filas:
        h.append(list(fila))
    salida = io.BytesIO()
    libro.save(salida)
    return SimpleUploadedFile("datos.xlsx", salida.getvalue())


class ImportadorBase(Base):
    def setUp(self):
        super().setUp()
        self.tec = self.usuario("tec", roles.TECNICO, unidades=[self.riego])
        self.client.force_login(self.tec)

    def subir(self, tabla, archivo, crear_catalogos=False):
        datos = {"tabla": tabla, "archivo": archivo}
        if crear_catalogos:
            datos["crear_catalogos"] = "1"
        self.client.post("/importar/tabular/", datos)
        return Importacion.objects.latest("id")

    def confirmar(self, imp):
        self.client.post(f"/importar/{imp.pk}/confirmar/")
        imp.refresh_from_db()
        return imp


class PlantillaTests(ImportadorBase):
    def test_plantilla_con_listas(self):
        p = self.proyecto(municipios=[self.m1])
        r = self.client.get("/importar/plantillas/mediciones.xlsx")
        self.assertEqual(r.status_code, 200)
        libro = load_workbook(io.BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ["Instrucciones", "Datos", "Listas"])
        encabezados = [c.value for c in libro["Datos"][1]]
        self.assertIn("Código de proyecto*", encabezados)
        self.assertIn("Desagregación", encabezados)  # opcional: sin *
        listas = [c.value for c in libro["Listas"]["A"]]
        self.assertEqual(listas[:2], ["Código de proyecto", p.codigo])
        self.assertTrue(libro["Datos"].data_validations.dataValidation, "listas desplegables")

    def test_encargado_no_ve_plantillas_que_no_puede_cargar(self):
        self.client.force_login(self.usuario("enc", roles.ENCARGADO_MUNICIPAL, municipios=[self.m1]))
        self.assertEqual(self.client.get("/importar/plantillas/financiamientos.xlsx").status_code, 404)
        self.assertEqual(self.client.get("/importar/plantillas/mediciones.xlsx").status_code, 200)


class TabularTests(ImportadorBase):
    def test_validar_confirmar_y_revertir(self):
        existente = self.proyecto(nombre="Viejo nombre", gestion=2025)
        archivo = xlsx(["Código", "Unidad*", "Nombre*", "Gestión", "Orden de proceder"],
                       ["", "riego", "Proyecto nuevo", 2026, "15/03/2026"],
                       [existente.codigo, "Riego", "Nombre corregido", 2025, date(2025, 5, 1)])
        imp = self.subir("proyectos", archivo)
        self.assertEqual(imp.estado, Importacion.VALIDADA, imp.resumen)
        self.assertEqual((imp.resumen["creados"], imp.resumen["actualizados"]), (1, 1))
        self.assertFalse(Proyecto.objects.filter(nombre="Proyecto nuevo").exists(), "validar no guarda")

        imp = self.confirmar(imp)
        self.assertEqual(imp.estado, Importacion.APLICADA)
        nuevo = Proyecto.objects.get(nombre="Proyecto nuevo")
        self.assertTrue(nuevo.codigo.startswith("RIE-2026-"))
        self.assertEqual(nuevo.fecha_orden_proceder, date(2026, 3, 15))
        existente.refresh_from_db()
        self.assertEqual(existente.nombre, "Nombre corregido")

        self.client.post(f"/importar/{imp.pk}/revertir/")
        imp.refresh_from_db()
        self.assertEqual(imp.estado, Importacion.REVERTIDA)
        self.assertFalse(Proyecto.objects.filter(pk=nuevo.pk).exists())
        existente.refresh_from_db()
        self.assertEqual(existente.nombre, "Viejo nombre")

    def test_errores_por_fila_y_nada_se_guarda(self):
        p = self.proyecto()
        archivo = xlsx(["Código de proyecto", "Municipio"], [p.codigo, "Punata"], [p.codigo, "Narnia"], ["XXX-1", "Tacachi"])
        imp = self.subir("proyecto_municipios", archivo)
        self.assertEqual(imp.estado, Importacion.CON_ERRORES)
        self.assertEqual([e["fila"] for e in imp.resumen["errores"]], [3, 4])
        self.assertIn("Narnia", imp.resumen["errores"][0]["mensajes"][0])
        self.assertFalse(p.municipios.exists(), "con errores no se guarda ni la fila buena")
        libro = load_workbook(io.BytesIO(self.client.get(f"/importar/{imp.pk}/errores.xlsx").content))
        self.assertEqual(libro.active["C1"].value, "Errores")
        self.assertIn("Narnia", libro.active["C3"].value)

    def test_columna_obligatoria_faltante(self):
        imp = self.subir("financiamientos", xlsx(["Código de proyecto", "Entidad"], ["X", "GADC"]))
        self.assertEqual(imp.estado, Importacion.CON_ERRORES)
        self.assertIn("Monto comprometido", imp.resumen["errores"][0]["mensajes"][0])

    def test_catalogo_nuevo_solo_si_se_pide(self):
        p = self.proyecto()
        TipoEvento.objects.create(nombre="Paralización")
        archivo = lambda: xlsx(["Código de proyecto", "Fecha", "Tipo de evento", "Descripción"],
                               [p.codigo, "01/02/2026", "Paralizacion temporal", "Lluvias"])
        imp = self.subir("eventos", archivo())
        self.assertEqual(imp.estado, Importacion.CON_ERRORES)
        self.assertIn("¿Quiso decir: Paralización?", imp.resumen["errores"][0]["mensajes"][0])
        imp = self.confirmar(self.subir("eventos", archivo(), crear_catalogos=True))
        self.assertEqual(imp.estado, Importacion.APLICADA)
        self.assertTrue(TipoEvento.objects.filter(nombre="Paralizacion temporal").exists())
        self.client.post(f"/importar/{imp.pk}/revertir/")
        self.assertFalse(TipoEvento.objects.filter(nombre="Paralizacion temporal").exists(), "revertir quita el catálogo creado")
        self.assertFalse(Evento.objects.exists())

    def test_csv_con_punto_y_coma_y_coma_decimal(self):
        p = self.proyecto()
        csv = f"Código de proyecto;Entidad;Monto comprometido (Bs)\n{p.codigo};GADC;1.250.000,50\n".encode("cp1252")
        imp = self.confirmar(self.subir("financiamientos", SimpleUploadedFile("f.csv", csv)))
        self.assertEqual(imp.estado, Importacion.APLICADA, imp.resumen)
        self.assertEqual(Financiamiento.objects.get().monto_comprometido, Decimal("1250000.50"))

    def test_desembolso_pide_financiamiento_y_no_duplica(self):
        p = self.proyecto()
        archivo = lambda: xlsx(["Código de proyecto", "Entidad", "Fecha", "Monto (Bs)", "Comprobante"],
                               [p.codigo, "GADC", "10/04/2026", "100", "C-1"])
        imp = self.subir("desembolsos", archivo())
        self.assertIn("no tiene financiamiento", imp.resumen["errores"][0]["mensajes"][0])
        Financiamiento.objects.create(proyecto=p, entidad=self.gadc, monto_comprometido=Decimal("1000"))
        self.confirmar(self.subir("desembolsos", archivo()))
        imp = self.subir("desembolsos", archivo())
        self.assertEqual(imp.resumen["sin_cambios"], 1)
        self.assertEqual(Desembolso.objects.count(), 1)

    def test_geometrias_por_coordenadas_utm_y_wkt(self):
        p = self.proyecto()
        archivo = xlsx(["Código de proyecto", "Tipo de elemento", "Sistema de coordenadas",
                        "Coordenada 1 (latitud o este)", "Coordenada 2 (longitud o norte)", "Geometría WKT"],
                       [p.codigo, "Toma de agua", "WGS84 (latitud, longitud)", 0.5, 0.5, ""],
                       [p.codigo, "Canal", "", "", "", "LINESTRING(0.2 0.2, 1.5 0.5)"])
        imp = self.confirmar(self.subir("geometrias", archivo))
        self.assertEqual(imp.estado, Importacion.APLICADA, imp.resumen)
        self.assertEqual(set(p.geometrias.values_list("tipo_elemento__nombre", flat=True)), {"Toma de agua", "Canal"})
        self.assertEqual(set(p.municipios.all()), {self.m1, self.m2}, "municipios desde las geometrías")
        punto = p.geometrias.get(tipo_elemento=self.toma).geom
        self.assertAlmostEqual((punto.x, punto.y), (0.5, 0.5))


class DetallesTests(ImportadorBase):
    def test_detalle_de_residuos_y_forestacion(self):
        from proyectos.models import DetalleForestacion, DetalleResiduos, Proposito

        botadero = self.proyecto(unidad=self.residuos, nombre="Botadero X")
        self.tec.unidades.add(self.residuos)
        imp = self.confirmar(self.subir("detalle_residuos", xlsx(
            ["Código de proyecto", "Botadero", "Licencia", "Riesgo", "Toneladas urbano", "Toneladas rural"],
            [botadero.codigo, "Botadero Municipal X", "Sin Licencia", "RIESGO MEDIO", "1,5", "8,282"])))
        self.assertEqual(imp.estado, Importacion.APLICADA, imp.resumen)
        self.assertEqual(DetalleResiduos.objects.get().ton_total, Decimal("9.782"))
        imp = self.subir("detalle_residuos", xlsx(["Código de proyecto", "Riesgo"], [botadero.codigo, "RIESGO ALTO"]))
        self.assertEqual(imp.resumen["actualizados"], 1, "segunda carga actualiza, no duplica")

        Proposito.objects.create(nombre="Comercial")
        forestal = self.proyecto(nombre="Forestación 2025-2026 – Punata")
        imp = self.confirmar(self.subir("detalle_forestacion", xlsx(
            ["Código de proyecto", "Propósito", "Campaña", "Plantines"], [forestal.codigo, "comercial", "2025-2026", 25800])))
        self.assertEqual(imp.estado, Importacion.APLICADA, imp.resumen)
        self.assertEqual(DetalleForestacion.objects.get().plantines, 25800)


class AlcanceTests(ImportadorBase):
    def test_encargado_solo_en_su_municipio(self):
        propio, ajeno = self.proyecto(municipios=[self.m1]), self.proyecto(municipios=[self.m3])
        Indicador.objects.create(nombre="Familias beneficiadas", unidad_medida="familias")
        self.client.force_login(self.usuario("enc", roles.ENCARGADO_MUNICIPAL, municipios=[self.m1]))
        cab = ["Código de proyecto", "Indicador", "Momento", "Fecha", "Valor"]
        imp = self.subir("mediciones", xlsx(cab, [ajeno.codigo, "Familias beneficiadas", "Seguimiento", "01/10/2026", 40]))
        self.assertIn("fuera de su alcance", imp.resumen["errores"][0]["mensajes"][0])
        imp = self.confirmar(self.subir("mediciones", xlsx(cab, [propio.codigo, "familias BENEFICIADAS", "seguimiento", "01/10/2026", 40])))
        self.assertEqual(imp.estado, Importacion.APLICADA, imp.resumen)
        self.assertEqual(Medicion.objects.get().registrado_por.username, "enc")
        r = self.client.post("/importar/tabular/", {"tabla": "financiamientos", "archivo": xlsx(["x"], ["y"])})
        self.assertEqual(Importacion.objects.filter(destino="financiamientos").count(), 0, "tabla no permitida")
        self.assertEqual(r.status_code, 302)

    def test_no_ve_importaciones_ajenas(self):
        imp = self.subir("proyectos", xlsx(["Unidad", "Nombre"], ["Riego", "X"]))
        self.client.force_login(self.usuario("otro", roles.TECNICO, unidades=[self.riego]))
        self.assertEqual(self.client.get(f"/importar/{imp.pk}/").status_code, 404)


class ExportacionTests(ImportadorBase):
    def test_excel_sin_ids_y_reimportable(self):
        p = self.proyecto(nombre="Riego Punata", municipios=[self.m1, self.m2], gestion=2026)
        q = self.proyecto(nombre="Riego Lejano", municipios=[self.m3])
        Financiamiento.objects.create(proyecto=p, entidad=self.gam, monto_comprometido=Decimal("500"))
        r = self.client.get("/importar/exportar/excel/", {"municipio": self.m1.id})
        libro = load_workbook(io.BytesIO(r.content))
        self.assertIn("Proyectos", libro.sheetnames)
        filas = list(libro["Proyectos"].iter_rows(values_only=True))
        self.assertEqual([f[0] for f in filas[1:]], [p.codigo], "respeta el filtro de municipio")
        self.assertEqual(filas[1][1], "Riego", "unidad por nombre, no por id")
        self.assertEqual(list(libro["Financiamientos"].iter_rows(values_only=True))[1][:3], (p.codigo, "GAM Punata", 500))
        self.assertEqual(len(list(libro["Municipios de proyectos"].iter_rows())), 3, "encabezado + 2 municipios")

        salida = io.BytesIO()
        libro.save(salida)
        imp = self.subir("proyectos", SimpleUploadedFile("exportado.xlsx", salida.getvalue()))
        self.assertEqual(imp.estado, Importacion.VALIDADA, imp.resumen)
        self.assertEqual((imp.resumen["sin_cambios"], imp.resumen["creados"]), (1, 0), "ida y vuelta sin cambios")
        self.assertEqual(self.client.get("/importar/exportar/excel/", {"ids": str(q.pk)}).status_code, 200)

    def test_geojson_y_kml(self):
        p = self.proyecto(municipios=[self.m1])
        Geometria.objects.create(proyecto=p, tipo_elemento=self.toma, nombre="Toma 1", geom=Point(0.5, 0.5, srid=4326))
        datos = json.loads(self.client.get("/importar/exportar/geojson/").content)
        self.assertEqual(datos["features"][0]["properties"]["codigo"], p.codigo)
        kml = self.client.get("/importar/exportar/kml/").content.decode()
        self.assertIn("<Placemark>", kml)
        self.assertIn(p.codigo, kml)


GEOJSON = {"type": "FeatureCollection", "features": [
    {"type": "Feature", "properties": {"nombre": "Toma A"}, "geometry": {"type": "Point", "coordinates": [0.5, 0.5]}},
    {"type": "Feature", "properties": {"nombre": "Área A"},
     "geometry": {"type": "MultiPolygon", "coordinates": [[[[1.2, 0.2], [1.8, 0.2], [1.8, 0.8], [1.2, 0.2]]]]}},
]}

KML = """<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>
<Placemark><name>Pozo norte</name><Point><coordinates>0.4,0.6,2550</coordinates></Point></Placemark>
</Document></kml>"""


class EspacialTests(ImportadorBase):
    def subir_mapa(self, nombre, contenido, destino="geometrias"):
        self.client.post("/importar/espacial/", {"destino": destino, "sistema": "4326",
                                                  "archivo": SimpleUploadedFile(nombre, contenido)})
        return Importacion.objects.latest("id")

    def opciones(self, imp, **datos):
        self.client.post(f"/importar/{imp.pk}/opciones/", datos)
        imp.refresh_from_db()
        return imp

    def test_geojson_a_geometrias_de_un_proyecto(self):
        p = self.proyecto()
        imp = self.subir_mapa("capa.geojson", json.dumps(GEOJSON).encode())
        self.assertEqual(imp.estado, Importacion.PENDIENTE)
        self.assertEqual(imp.resumen["formas"], {"Point": 1, "Polygon": 1}, "el multipolígono se separa en partes")
        self.assertEqual(self.client.get(f"/importar/{imp.pk}/").status_code, 200)
        self.assertEqual(len(json.loads(self.client.get(f"/importar/{imp.pk}/vista-previa.geojson").content)["features"]), 2)

        imp = self.opciones(imp, proyecto=p.pk, tipo_Point=self.toma.pk, atributo_nombre="nombre")
        self.assertEqual(imp.estado, Importacion.CON_ERRORES, "falta el tipo para los polígonos")
        imp = self.opciones(imp, proyecto=p.pk, tipo_Point=self.toma.pk, tipo_Polygon=self.area.pk, atributo_nombre="nombre")
        self.assertEqual(imp.estado, Importacion.VALIDADA, imp.resumen)
        self.assertFalse(Geometria.objects.exists())
        imp = self.confirmar(imp)
        self.assertEqual(sorted(p.geometrias.values_list("nombre", flat=True)), ["Toma A", "Área A"])
        self.assertEqual(set(p.municipios.all()), {self.m1, self.m2})
        self.client.post(f"/importar/{imp.pk}/revertir/")
        self.assertFalse(Geometria.objects.exists())
        self.assertFalse(p.municipios.exists())

    def test_kml_y_kmz(self):
        p = self.proyecto()
        kmz = io.BytesIO()
        with zipfile.ZipFile(kmz, "w") as z:
            z.writestr("doc.kml", KML)
        for nombre, contenido in (("puntos.kml", KML.encode()), ("puntos.kmz", kmz.getvalue())):
            imp = self.subir_mapa(nombre, contenido)
            self.assertEqual(imp.resumen["formas"], {"Point": 1}, nombre)
            imp = self.opciones(imp, proyecto=p.pk, tipo_Point=self.toma.pk, atributo_nombre="Name")
            self.assertEqual(imp.estado, Importacion.VALIDADA, imp.resumen)
        imp = self.confirmar(imp)
        g = p.geometrias.get()
        self.assertEqual((g.nombre, round(g.geom.x, 3), round(g.geom.y, 3)), ("Pozo norte", 0.4, 0.6), "sin altura")

    def test_zip_con_rutas_peligrosas(self):
        malo = io.BytesIO()
        with zipfile.ZipFile(malo, "w") as z:
            z.writestr("../fuera.kml", KML)
        imp = self.subir_mapa("malo.zip", malo.getvalue())
        self.assertEqual(imp.estado, Importacion.CON_ERRORES)

    def test_actualizar_limites_y_revertir(self):
        admin = Usuario.objects.create_superuser("root", password="clave-de-prueba-123")
        self.client.force_login(admin)
        original = Municipio.objects.get(pk=self.m1.pk).geom
        nuevo = {"type": "FeatureCollection", "features": [{"type": "Feature", "properties": {"COD": "031401"},
                 "geometry": {"type": "Polygon", "coordinates": [[[0, 0], [1.1, 0], [1.1, 1], [0, 1], [0, 0]]]}}]}
        imp = self.subir_mapa("limites.geojson", json.dumps(nuevo).encode(), destino="limites")
        imp = self.opciones(imp, atributo_municipio="COD")
        self.assertEqual(imp.estado, Importacion.VALIDADA, imp.resumen)
        imp = self.confirmar(imp)
        self.assertAlmostEqual(Municipio.objects.get(pk=self.m1.pk).geom.extent[2], 1.1)
        self.client.post(f"/importar/{imp.pk}/revertir/")
        self.assertTrue(Municipio.objects.get(pk=self.m1.pk).geom.equals(original))

    def test_tecnico_no_actualiza_limites(self):
        self.client.post("/importar/espacial/", {"destino": "limites", "sistema": "4326",
                                                  "archivo": SimpleUploadedFile("l.geojson", json.dumps(GEOJSON).encode())})
        self.assertFalse(Importacion.objects.filter(destino="limites").exists())


class ReversionTests(ImportadorBase):
    def test_no_revierte_si_se_modifico_despues(self):
        p = self.proyecto(nombre="A")
        imp = self.confirmar(self.subir("proyectos", xlsx(["Código", "Unidad", "Nombre"], [p.codigo, "Riego", "B"])))
        p.refresh_from_db()
        p.nombre = "C"
        p.save()
        self.client.post(f"/importar/{imp.pk}/revertir/")
        imp.refresh_from_db()
        self.assertEqual(imp.estado, Importacion.APLICADA)
        p.refresh_from_db()
        self.assertEqual(p.nombre, "C", "no pisa cambios posteriores")


class CargaInicialTests(ImportadorBase):
    def test_catalogos_iniciales_y_entidades(self):
        call_command("cargar_catalogos", stdout=io.StringIO())
        imp = self.subir("entidades", xlsx(["Nombre", "Tipo de fuente", "Municipio"], ["Cooperación X", "Externo", ""]))
        self.client.force_login(Usuario.objects.create_superuser("root", password="clave-de-prueba-123"))
        imp = self.confirmar(self.subir("entidades", xlsx(["Nombre", "Tipo de fuente", "Municipio"], ["Cooperación X", "Externo", ""])))
        self.assertEqual(imp.estado, Importacion.APLICADA, imp.resumen)
        self.assertTrue(Entidad.objects.filter(nombre="Cooperación X").exists())
